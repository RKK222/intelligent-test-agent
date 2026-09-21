#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Non-interactive BDSP result-query CLI for OpenCode skills."""

import argparse
import configparser
import json
import os
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
DEFAULT_CONFIG = BASE_DIR / "config" / "database.ini"
DEFAULT_KEY = BASE_DIR / "config" / ".encryption.key"
DRIVER_JAR = BASE_DIR / "lib" / "GaussDBV5-503.1.0.SPC1700_23.09.22.jar"
JAVA_SOURCE = Path(__file__).resolve().parent / "GaussDBQuery.java"
JAVA_CLASS = Path(__file__).resolve().parent / "GaussDBQuery.class"


def eprint(*args, **kwargs):
    print(*args, file=sys.stderr, **kwargs)


def validate_env(value):
    env = str(value).strip().upper()
    if env not in {"JC2", "JC4", "JC6"}:
        raise ValueError(f"不支持的调度环境: {value}; 仅支持 JC2/JC4/JC6")
    return env


def validate_date(value):
    if not value:
        return datetime.now().strftime("%Y%m%d")
    value = str(value).strip()
    try:
        datetime.strptime(value, "%Y%m%d")
    except ValueError:
        raise ValueError(f"调度日期必须为 YYYYMMDD: {value}")
    return value


def read_excel(path):
    path = Path(path).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"输入文件不存在: {path}")
    rows = []
    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("缺少 openpyxl，请先安装 requirements.txt") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        headers = [str(v).strip() if v is not None else "" for v in next(it)]
        for values in it:
            if not any(v is not None and str(v).strip() for v in values):
                continue
            rows.append({headers[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(values) if i < len(headers)})
        wb.close()
    elif path.suffix.lower() == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("缺少 xlrd，请先安装 requirements.txt") from exc
        book = xlrd.open_workbook(str(path))
        sheet = book.sheet_by_index(0)
        headers = [str(sheet.cell_value(0, c)).strip() for c in range(sheet.ncols)]
        for r in range(1, sheet.nrows):
            values = [sheet.cell_value(r, c) for c in range(sheet.ncols)]
            if not any(str(v).strip() for v in values):
                continue
            row = {}
            for c, v in enumerate(values):
                if isinstance(v, float) and v.is_integer():
                    v = int(v)
                row[headers[c]] = str(v).strip()
            rows.append(row)
    else:
        raise ValueError("只支持 .xls 或 .xlsx 文件")
    return rows


def decrypt_password(value, key_file):
    if not value:
        return value
    if value.startswith("env:"):
        env_name = value.split(":", 1)[1].strip()
        password = os.environ.get(env_name)
        if password is None:
            raise RuntimeError(f"环境变量 {env_name} 未设置")
        return password
    if not value.startswith("gAAAA"):
        return value
    try:
        from cryptography.fernet import Fernet
    except ImportError as exc:
        raise RuntimeError("缺少 cryptography，请先安装 requirements.txt") from exc
    key_path = Path(key_file)
    if not key_path.exists():
        raise FileNotFoundError(f"加密密钥不存在: {key_path}")
    cipher = Fernet(key_path.read_bytes())
    return cipher.decrypt(value.encode("utf-8")).decode("utf-8")


def load_config(path, key_file):
    cfg_path = Path(path).expanduser().resolve()
    if not cfg_path.exists():
        raise FileNotFoundError(f"数据库配置不存在: {cfg_path}")
    parser = configparser.ConfigParser()
    parser.read(cfg_path, encoding="utf-8")
    result = {}
    for env in ("JC2", "JC4", "JC6"):
        sec = env.lower()
        if sec not in parser:
            continue
        section = parser[sec]
        result[env] = {
            "host": section.get("host", "").strip(),
            "port": int(section.get("port", "0")),
            "database": section.get("database", "").strip(),
            "username": section.get("username", "").strip(),
            "password": decrypt_password(section.get("password", ""), key_file),
        }
    if not result:
        raise RuntimeError("数据库配置中没有 jc2/jc4/jc6")
    return result


def ensure_java_class():
    java = shutil.which("java")
    if not java:
        raise RuntimeError("未找到 java 命令")
    if JAVA_CLASS.exists():
        return java
    javac = shutil.which("javac")
    if not javac:
        raise RuntimeError("未找到 GaussDBQuery.class，且系统没有 javac 可用于编译")
    cmd = [javac, "-encoding", "UTF-8", "-source", "8", "-target", "8", str(JAVA_SOURCE)]
    proc = subprocess.run(cmd, cwd=str(JAVA_SOURCE.parent), capture_output=True, text=True)
    if proc.returncode != 0:
        raise RuntimeError(f"Java 编译失败: {proc.stderr.strip()}")
    return java


def java_query(db, sql):
    java = ensure_java_class()
    if not DRIVER_JAR.exists():
        raise FileNotFoundError(f"GaussDB JDBC 驱动不存在: {DRIVER_JAR}")
    cmd = [
        java,
        "-cp",
        os.pathsep.join([str(JAVA_SOURCE.parent), str(DRIVER_JAR)]),
        "GaussDBQuery",
        db["host"],
        str(db["port"]),
        db["database"],
        db["username"],
        sql,
    ]
    env = os.environ.copy()
    env["GAUSSDB_PASSWORD"] = db["password"]
    proc = subprocess.run(cmd, cwd=str(JAVA_SOURCE.parent), env=env, capture_output=True, text=True, timeout=120)
    if proc.returncode != 0:
        msg = proc.stderr.strip().splitlines()[-1] if proc.stderr.strip() else "Java 查询失败"
        raise RuntimeError(msg)
    try:
        data = json.loads(proc.stdout)
    except json.JSONDecodeError as exc:
        raise RuntimeError("Java 查询结果不是有效 JSON") from exc
    if "error" in data:
        raise RuntimeError(str(data["error"]))
    return data.get("rows", [])


def sql_literal(value):
    return str(value).replace("'", "''")


def build_sql(kind, row):
    group = sql_literal(row["作业组名"])
    if kind == "job":
        job = sql_literal(row["作业名"])
        return f"SELECT * FROM bdsp.etl_job AS a WHERE a.etl_system = '{group}' AND a.etl_job = '{job}'"
    return f"SELECT * FROM bdsp.etl_job AS a WHERE a.etl_system = '{group}'"


def export_xls(rows, path):
    try:
        import xlwt
    except ImportError as exc:
        raise RuntimeError("缺少 xlwt，请先安装 requirements.txt") from exc
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = xlwt.Workbook()
    sheet = book.add_sheet("执行结果")
    if not rows:
        sheet.write(0, 0, "无查询数据")
        book.save(str(path))
        return
    columns = list(rows[0].keys())
    header_style = xlwt.easyxf("font: bold on; borders: top thin, right thin, bottom thin, left thin;")
    data_style = xlwt.easyxf("borders: top thin, right thin, bottom thin, left thin;")
    for c, name in enumerate(columns):
        sheet.write(0, c, name, header_style)
    widths = [len(str(name)) for name in columns]
    for r, row in enumerate(rows, 1):
        for c, name in enumerate(columns):
            value = "" if row.get(name) is None else str(row.get(name))
            sheet.write(r, c, value, data_style)
            widths[c] = min(50, max(widths[c], len(value)))
    for c, width in enumerate(widths):
        sheet.col(c).width = max(1500, min(12000, (width + 2) * 256))
    book.save(str(path))


def direct_row(args, kind):
    row = {
        "调度环境": args.env,
        "调度应用": args.app,
        "作业组名": args.group,
        "调度日期": args.date or "",
    }
    if kind == "job":
        row["作业名"] = args.job
    return row


def validate_rows(rows, kind):
    required = ["调度环境", "调度应用", "作业组名"]
    if kind == "job":
        required.append("作业名")
    if not rows:
        raise ValueError("没有可处理的数据")
    normalized = []
    for i, raw in enumerate(rows, 1):
        missing = [k for k in required if not str(raw.get(k, "")).strip()]
        if missing:
            raise ValueError(f"第 {i} 行缺少字段: {', '.join(missing)}")
        row = dict(raw)
        row["调度环境"] = validate_env(raw["调度环境"])
        row["调度应用"] = str(raw["调度应用"]).strip().upper()
        row["作业组名"] = str(raw["作业组名"]).strip()
        row["调度日期"] = validate_date(raw.get("调度日期", ""))
        if kind == "job":
            row["作业名"] = str(raw["作业名"]).strip()
        normalized.append(row)
    return normalized


def build_parser():
    parser = argparse.ArgumentParser(description="BDSP 作业组/单作业执行结果非交互查询工具")
    sub = parser.add_subparsers(dest="kind", required=True)
    for kind in ("group", "job"):
        p = sub.add_parser(kind, help="group=作业组查询, job=单作业查询")
        p.add_argument("--input", help="批量 Excel (.xls/.xlsx)")
        p.add_argument("--env")
        p.add_argument("--app")
        p.add_argument("--group")
        if kind == "job":
            p.add_argument("--job")
        p.add_argument("--date")
        p.add_argument("--config", default=str(DEFAULT_CONFIG))
        p.add_argument("--key-file", default=str(DEFAULT_KEY))
        p.add_argument("--output-dir")
        p.add_argument("--validate-only", action="store_true")
    return parser


def main():
    args = build_parser().parse_args()
    kind = args.kind
    try:
        if args.input:
            rows = read_excel(args.input)
        else:
            required_args = [args.env, args.app, args.group]
            if kind == "job":
                required_args.append(args.job)
            if any(v is None for v in required_args):
                raise ValueError("未使用 --input 时必须提供 --env/--app/--group，单作业查询还需 --job")
            rows = [direct_row(args, kind)]
        rows = validate_rows(rows, kind)

        if args.validate_only:
            # Validate config structure without printing or connecting.
            load_config(args.config, args.key_file)
            print(json.dumps({"validated": True, "kind": kind, "count": len(rows)}, ensure_ascii=False))
            return 0

        db_configs = load_config(args.config, args.key_file)
        output_dir = Path(args.output_dir).expanduser().resolve() if args.output_dir else Path.cwd() / f"执行结果_{datetime.now().strftime('%Y%m%d')}"
        success = 0
        failed = 0
        details = []
        for i, row in enumerate(rows, 1):
            env = row["调度环境"]
            try:
                if env not in db_configs:
                    raise RuntimeError(f"数据库配置中未配置 {env}")
                sql = build_sql(kind, row)
                data = java_query(db_configs[env], sql)
                app = row["调度应用"]
                group = row["作业组名"]
                date = row["调度日期"]
                if kind == "job":
                    filename = f"{app}_{group}_{row['作业名']}_{date}_{env}.xls"
                else:
                    filename = f"{app}_{group}_{date}_{env}.xls"
                target = output_dir / env / filename
                export_xls(data, target)
                success += 1
                details.append({"index": i, "ok": True, "records": len(data), "file": str(target)})
            except Exception as exc:
                failed += 1
                details.append({"index": i, "ok": False, "error": str(exc)})
        print(json.dumps({"details": details}, ensure_ascii=False, indent=2))
        print(json.dumps({"total": len(rows), "success": success, "failed": failed, "output_dir": str(output_dir)}, ensure_ascii=False))
        return 0 if failed == 0 else 1
    except Exception as exc:
        eprint(f"ERROR: {exc}")
        print(json.dumps({"total": 0, "success": 0, "failed": 1, "error": str(exc)}, ensure_ascii=False))
        return 2


if __name__ == "__main__":
    sys.exit(main())
