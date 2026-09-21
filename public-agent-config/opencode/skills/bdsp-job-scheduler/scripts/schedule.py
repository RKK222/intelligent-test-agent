#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Non-interactive BDSP scheduler CLI for OpenCode skills."""

import argparse
import copy
import json
import os
import platform
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
CONFIG_DIR = BASE_DIR / "config"
STATION_DIR = BASE_DIR / "station"
ENV_MAPPING = {
    "JC2": "BDSP_JC2",
    "JC4": "BDSP_JC4",
    "JC6": "BDSP_JC6",
    "BDSP_JC2": "BDSP_JC2",
    "BDSP_JC4": "BDSP_JC4",
    "BDSP_JC6": "BDSP_JC6",
}


def eprint(*args, **kwargs):
    print(*args, file=sys.stderr, **kwargs)


def validate_date(value):
    try:
        datetime.strptime(value, "%Y%m%d")
    except ValueError:
        raise ValueError(f"调度日期必须为 YYYYMMDD: {value}")
    return value


def standardize_env(value):
    key = str(value).strip().upper()
    if key not in ENV_MAPPING:
        raise ValueError(f"不支持的调度环境: {value}; 仅支持 JC2/JC4/JC6")
    return ENV_MAPPING[key]


def read_excel(path):
    path = Path(path).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"输入文件不存在: {path}")
    suffix = path.suffix.lower()
    rows = []
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("缺少 openpyxl，请先安装 requirements.txt") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        values = ws.iter_rows(values_only=True)
        headers = [str(v).strip() if v is not None else "" for v in next(values)]
        for values_row in values:
            if not any(v is not None and str(v).strip() for v in values_row):
                continue
            rows.append({headers[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(values_row) if i < len(headers)})
        wb.close()
    elif suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("缺少 xlrd，请先安装 requirements.txt") from exc
        book = xlrd.open_workbook(str(path))
        sheet = book.sheet_by_index(0)
        headers = [str(sheet.cell_value(0, c)).strip() for c in range(sheet.ncols)]
        for r in range(1, sheet.nrows):
            values_row = [sheet.cell_value(r, c) for c in range(sheet.ncols)]
            if not any(str(v).strip() for v in values_row):
                continue
            row = {}
            for c, v in enumerate(values_row):
                if c >= len(headers):
                    break
                if isinstance(v, float) and v.is_integer():
                    v = int(v)
                row[headers[c]] = str(v).strip()
            rows.append(row)
    else:
        raise ValueError("只支持 .xls 或 .xlsx 文件")
    return rows


def load_template(kind):
    name = "request_template.json" if kind == "group" else "request_template_job.json"
    path = CONFIG_DIR / name
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def find_capture_binary():
    system = platform.system()
    machine = platform.machine().lower()
    if system == "Windows":
        candidates = ["capture-proxy-windows-amd64.exe", "capture-proxy-windows-386.exe"]
    elif system == "Darwin":
        if machine in {"arm64", "aarch64"}:
            candidates = ["capture-proxy-macos-arm64", "capture-proxy-mac"]
        else:
            candidates = ["capture-proxy-mac"]
    else:
        if machine in {"arm64", "aarch64"}:
            candidates = ["capture-proxy-linux-arm64"]
        else:
            candidates = ["capture-proxy-linux-amd64"]
    for name in candidates:
        p = STATION_DIR / name
        if p.exists():
            return p
    return None


def strip_bearer(value):
    value = str(value).strip()
    prefix = "Bearer "
    return value[len(prefix):].strip() if value.startswith(prefix) else value


def load_auth_from_file(path):
    with Path(path).expanduser().open("r", encoding="utf-8") as f:
        data = json.load(f)
    token = data.get("token") or data.get("Authorization")
    if not token:
        raise RuntimeError("授权文件中没有 token/Authorization 字段")
    return strip_bearer(token)


def obtain_auth_token(auth_file=None):
    env_token = os.environ.get("BDSP_AUTH_TOKEN", "").strip()
    if env_token:
        return strip_bearer(env_token)
    if auth_file:
        return load_auth_from_file(auth_file)

    result_file = STATION_DIR / "result.txt"
    if result_file.exists():
        try:
            result_file.unlink()
        except OSError:
            pass

    binary = find_capture_binary()
    if not binary:
        raise RuntimeError("未找到适配当前系统/架构的授权捕获程序")
    if platform.system() != "Windows":
        binary.chmod(binary.stat().st_mode | 0o111)

    eprint(f"正在运行授权捕获程序: {binary.name}；请在自动打开的浏览器中完成授权页面访问")
    log_file = STATION_DIR / "capture-proxy.log"
    try:
        # 新版 station 会把 Authorization 打印到自身终端/日志。OpenCode 调用时静默其
        # stdout/stderr，避免令牌进入 Agent 上下文；授权结果仅从 result.txt 读取。
        proc = subprocess.run(
            [str(binary)],
            cwd=str(STATION_DIR),
            check=False,
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        if proc.returncode != 0:
            raise RuntimeError(f"授权捕获程序退出码: {proc.returncode}")
        if not result_file.exists():
            raise RuntimeError("授权捕获完成但未生成 station/result.txt")
        token = load_auth_from_file(result_file)
        eprint("授权捕获完成，继续执行调度请求")
        return token
    finally:
        # result.txt / capture-proxy.log 都可能包含 Authorization，不作为 Skill 运行资产保留。
        for sensitive_file in (result_file, log_file):
            try:
                sensitive_file.unlink(missing_ok=True)
            except OSError:
                pass


def render_request(kind, row, auth_token):
    template = copy.deepcopy(load_template(kind))
    env = standardize_env(row["调度环境"])
    date = validate_date(row["调度日期"])
    user_id = str(row["统一认证号"]).strip()
    app = str(row["调度应用"]).strip().upper()
    group = str(row["作业组名"]).strip()
    job = str(row.get("作业名", "")).strip()
    if not all([user_id, app, group]):
        raise ValueError("统一认证号、调度应用、作业组名不能为空")
    if kind == "job" and not job:
        raise ValueError("单作业调度时作业名不能为空")
    base_url = os.environ.get("BDSP_BASE_URL", "").strip().rstrip("/")
    if not base_url:
        raise ValueError("未设置 BDSP_BASE_URL，不能发送调度请求")

    mapping = {
        "BDSP_BASE_URL": base_url,
        "AUTH_TOKEN": auth_token or "<AUTH_TOKEN>",
        "USERINFO": os.environ.get("BDSP_USERINFO", "").strip(),
        "ENV": env,
        "APP_NAME": app,
        "APP_GROUP": group,
        "ETL_SYSTEM": group,
        "JOB_NAME": job,
        "USER_ID": user_id,
        "BATCH_DATE": date,
    }
    for key, value in list(template.get("headers", {}).items()):
        if isinstance(value, str):
            rendered = value.format(**mapping)
            if key.lower() == "authorization" and rendered and not rendered.startswith("Bearer "):
                rendered = "Bearer " + rendered
            template["headers"][key] = rendered
    if isinstance(template.get("url"), str):
        template["url"] = template["url"].format(**mapping)
    for key, value in list(template.get("params", {}).items()):
        if isinstance(value, str):
            template["params"][key] = value.format(**mapping)
    return template


def send_request(request_data, retries):
    try:
        import requests
    except ImportError as exc:
        raise RuntimeError("缺少 requests，请先安装 requirements.txt") from exc

    last_error = None
    for attempt in range(1, retries + 1):
        try:
            response = requests.post(
                request_data["url"],
                json=request_data["params"],
                headers=request_data["headers"],
                timeout=30,
            )
            try:
                body = response.json()
            except Exception:
                body = {"code": response.status_code, "msg": response.text[:500]}
            ok = body.get("code") == 20000 and body.get("msg") == "成功"
            return ok, body
        except Exception as exc:
            last_error = str(exc)
            if attempt < retries:
                time.sleep(3)
    return False, {"code": -1, "msg": last_error or "请求失败"}


def row_from_args(args, kind):
    row = {
        "统一认证号": args.user_id,
        "调度环境": args.env,
        "调度应用": args.app,
        "作业组名": args.group,
        "调度日期": args.date,
    }
    if kind == "job":
        row["作业名"] = args.job
    return row


def validate_rows(rows, kind):
    required = ["统一认证号", "调度环境", "调度应用", "作业组名", "调度日期"]
    if kind == "job":
        required.append("作业名")
    if not rows:
        raise ValueError("没有可处理的数据")
    for index, row in enumerate(rows, 1):
        missing = [key for key in required if not str(row.get(key, "")).strip()]
        if missing:
            raise ValueError(f"第 {index} 行缺少字段: {', '.join(missing)}")
        standardize_env(row["调度环境"])
        validate_date(str(row["调度日期"]).strip())


def build_parser():
    parser = argparse.ArgumentParser(description="BDSP 作业组/单作业非交互调度工具")
    sub = parser.add_subparsers(dest="kind", required=True)
    for kind in ("group", "job"):
        p = sub.add_parser(kind, help="group=作业组调度, job=单作业调度")
        p.add_argument("--input", help="批量 Excel (.xls/.xlsx)")
        p.add_argument("--user-id")
        p.add_argument("--env")
        p.add_argument("--app")
        p.add_argument("--group")
        if kind == "job":
            p.add_argument("--job")
        p.add_argument("--date")
        p.add_argument("--auth-file", help="已有授权 JSON 文件")
        p.add_argument("--retries", type=int, default=1)
        p.add_argument("--execute", action="store_true", help="真正发送调度请求；缺省仅预览")
    return parser


def main():
    args = build_parser().parse_args()
    kind = args.kind
    if args.input:
        rows = read_excel(args.input)
    else:
        needed = [args.user_id, args.env, args.app, args.group, args.date]
        if kind == "job":
            needed.append(args.job)
        if any(v is None for v in needed):
            raise SystemExit("未使用 --input 时必须提供完整直接参数；运行 --help 查看参数")
        rows = [row_from_args(args, kind)]

    try:
        validate_rows(rows, kind)
        if not args.execute:
            for row in rows:
                render_request(kind, row, None)
            eprint(f"参数校验通过，共 {len(rows)} 条记录；添加 --execute 可正式执行调度")
            print(json.dumps({"total": len(rows), "success": 0, "failed": 0, "executed": False}, ensure_ascii=False))
            return 0

        auth_token = obtain_auth_token(args.auth_file)
        success = 0
        failed = 0
        details = []
        for i, row in enumerate(rows, 1):
            request_data = render_request(kind, row, auth_token)
            ok, body = send_request(request_data, max(1, args.retries))
            details.append({"index": i, "ok": ok, "code": body.get("code"), "msg": body.get("msg")})
            if ok:
                success += 1
            else:
                failed += 1
        print(json.dumps({"details": details}, ensure_ascii=False, indent=2))
        print(json.dumps({"total": len(rows), "success": success, "failed": failed, "executed": True}, ensure_ascii=False))
        return 0 if failed == 0 else 1
    except Exception as exc:
        eprint(f"ERROR: {exc}")
        print(json.dumps({"total": len(rows) if 'rows' in locals() else 0, "success": 0, "failed": len(rows) if 'rows' in locals() else 1, "error": str(exc)}, ensure_ascii=False))
        return 2


if __name__ == "__main__":
    sys.exit(main())
